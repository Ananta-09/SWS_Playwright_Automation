from openpyxl import load_workbook

def read_excel(file_path, sheet_name):
    workbook = load_workbook(file_path)
    worksheet = workbook[sheet_name]
    data = []
    for username,expected_result in worksheet.iter_rows(min_row=2,max_col=2,values_only=True):
        data.append((username,expected_result))
    workbook.close()
    return data
    